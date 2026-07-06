from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha1
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import pandas as pd
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..utils import ensure_dir, now_iso, slugify, write_json
from .scoring import RankingResult, add_comparable_keys


@dataclass(frozen=True)
class IdeaChartPackageResult:
    """Structured output from idea chart package generation."""

    output_dir: Path
    buy_candidates_path: Path
    sell_candidates_path: Path
    chart_metadata_path: Path
    package_metadata_path: Path
    chart_metadata: pd.DataFrame
    metadata: dict[str, Any]
    readable_summary_path: Path | None = None
    readable_workbook_path: Path | None = None


def _candidate_subset(candidates: pd.DataFrame, top_n: int | None) -> pd.DataFrame:
    if top_n is None:
        return candidates.copy().reset_index(drop=True)
    return candidates.head(max(int(top_n), 0)).copy().reset_index(drop=True)


def _candidate_filename(side: str, rank: int, candidate: pd.Series) -> str:
    raw_name = str(candidate.get("raw_story_name") or candidate.get("comparable_key") or "idea")
    comparable_key = str(candidate.get("comparable_key") or raw_name)
    digest = sha1(comparable_key.encode("utf-8")).hexdigest()[:8]
    clean_name = slugify(raw_name)[:70] or "idea"
    return f"{rank:03d}_{side}_{clean_name}_{digest}.png"


def _series_for_candidate(
    normalized_frame: pd.DataFrame,
    candidate: pd.Series,
    *,
    include_full_history: bool,
) -> pd.DataFrame:
    key = str(candidate.get("comparable_key") or "")
    if not key:
        return pd.DataFrame(columns=normalized_frame.columns)

    series = normalized_frame[normalized_frame["comparable_key"] == key].copy()
    if not include_full_history and "is_default_focus_window" in series.columns:
        series = series[series["is_default_focus_window"].fillna(False)].copy()

    series["date"] = pd.to_datetime(series["date"], errors="coerce")
    series["value"] = pd.to_numeric(series["value"], errors="coerce")
    return series.dropna(subset=["date", "value"]).sort_values("date").reset_index(drop=True)


def _chart_title(side: str, rank: int, candidate: pd.Series) -> str:
    raw_name = str(candidate.get("raw_story_name") or "Unknown Story")
    z_score = candidate.get("z_score")
    current = candidate.get("current_value")
    try:
        z_label = f"{float(z_score):+.2f}z"
    except (TypeError, ValueError):
        z_label = "z n/a"
    try:
        current_label = f"{float(current):.4f}"
    except (TypeError, ValueError):
        current_label = "n/a"
    return f"{side.upper()} #{rank}: {raw_name} | latest {current_label} | {z_label}"


def _save_idea_chart(
    series: pd.DataFrame,
    candidate: pd.Series,
    output_path: Path,
    *,
    side: str,
    rank: int,
    value_label: str,
) -> None:
    fig, ax = plt.subplots(figsize=(12, 6))
    ax.plot(series["date"], series["value"], color="#1f6feb", linewidth=2.2, marker="o", markersize=3.5)

    mean_value = candidate.get("mean")
    low_value = candidate.get("low")
    high_value = candidate.get("high")
    latest_date = series["date"].max()
    latest_value = float(series.loc[series["date"] == latest_date, "value"].mean())

    if pd.notna(mean_value):
        ax.axhline(float(mean_value), color="#6b7280", linewidth=1.4, linestyle="--", label="Mean")
    if pd.notna(low_value):
        ax.axhline(float(low_value), color="#16a34a", linewidth=1.0, linestyle=":", label="Low")
    if pd.notna(high_value):
        ax.axhline(float(high_value), color="#dc2626", linewidth=1.0, linestyle=":", label="High")

    marker_color = "#16a34a" if side == "buy" else "#dc2626"
    ax.scatter([latest_date], [latest_value], color=marker_color, s=70, zorder=5, label="Latest")
    ax.set_title(_chart_title(side, rank, candidate), fontsize=12, fontweight="bold")
    ax.set_xlabel("Date")
    ax.set_ylabel(value_label)
    ax.grid(alpha=0.25)
    ax.legend(loc="best", fontsize=8)
    fig.autofmt_xdate()
    fig.tight_layout()
    fig.savefig(output_path, dpi=180, bbox_inches="tight")
    plt.close(fig)


def _build_chart_records(
    normalized_frame: pd.DataFrame,
    candidates: pd.DataFrame,
    output_dir: Path,
    *,
    side: str,
    include_full_history: bool,
    value_label: str,
) -> list[dict[str, Any]]:
    chart_dir = ensure_dir(output_dir / "charts" / side)
    records: list[dict[str, Any]] = []

    for idx, candidate in candidates.iterrows():
        rank = int(idx) + 1
        series = _series_for_candidate(
            normalized_frame,
            candidate,
            include_full_history=include_full_history,
        )
        if series.empty:
            records.append(
                {
                    "side": side,
                    "rank": rank,
                    "comparable_key": candidate.get("comparable_key"),
                    "raw_story_name": candidate.get("raw_story_name"),
                    "chart_path": None,
                    "status": "no_series_data",
                }
            )
            continue

        filename = _candidate_filename(side, rank, candidate)
        chart_path = chart_dir / filename
        _save_idea_chart(
            series,
            candidate,
            chart_path,
            side=side,
            rank=rank,
            value_label=value_label,
        )
        records.append(
            {
                "side": side,
                "rank": rank,
                "comparable_key": candidate.get("comparable_key"),
                "raw_story_name": candidate.get("raw_story_name"),
                "chart_path": str(chart_path),
                "status": "ok",
                "latest_date": str(pd.to_datetime(series["date"].max()).date()),
                "series_points": int(len(series)),
                "current_value": candidate.get("current_value"),
                "z_score": candidate.get("z_score"),
                "mean": candidate.get("mean"),
                "low": candidate.get("low"),
                "high": candidate.get("high"),
                "observations": candidate.get("observations"),
            }
        )

    return records


def _relative_chart_path(chart_path: Any, package_dir: Path) -> str:
    if chart_path in (None, "") or pd.isna(chart_path):
        return ""
    path = Path(str(chart_path))
    try:
        return path.relative_to(package_dir).as_posix()
    except ValueError:
        return str(path)


def _friendly_value(value: Any) -> Any:
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    if isinstance(value, pd.Timestamp):
        return value.date().isoformat()
    return value


def _friendly_year(value: Any) -> Any:
    value = _friendly_value(value)
    if value is None:
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return value


def _friendly_coupon(value: Any) -> str | None:
    value = _friendly_value(value)
    if value is None:
        return None
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return str(value)
    if numeric.is_integer():
        return f"{numeric:.1f}"
    return f"{numeric:.3f}".rstrip("0").rstrip(".")


def _story_attribute(candidate: pd.Series) -> str:
    """Return the single client-facing story descriptor used in report sheets."""

    program = str(_friendly_value(candidate.get("program")) or "").strip()
    values = [
        _friendly_value(candidate.get("geography")),
        _friendly_value(candidate.get("loan_balance_bucket")),
        _friendly_value(candidate.get("story_tag")),
    ]
    family = str(_friendly_value(candidate.get("story_family")) or "").strip()
    if family and family.upper() != program.upper():
        values.append(family)

    parts: list[str] = []
    for value in values:
        text = str(value or "").strip()
        if text and text not in parts:
            parts.append(text)
    return " / ".join(parts) if parts else "Base story"


def _candidate_summary_rows(
    candidates: pd.DataFrame,
    chart_lookup: dict[tuple[str, int], pd.Series],
    package_dir: Path,
    *,
    side: str,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for idx, candidate in candidates.reset_index(drop=True).iterrows():
        rank = int(idx) + 1
        chart = chart_lookup.get((side, rank))
        chart_path = chart.get("chart_path") if chart is not None else None
        rows.append(
            {
                "Side": side.upper(),
                "Rank": rank,
                "Story Name": _friendly_value(candidate.get("raw_story_name")),
                "Program": _friendly_value(candidate.get("program")),
                "Vintage": _friendly_year(candidate.get("origination_year")),
                "Coupon": _friendly_coupon(candidate.get("coupon")),
                "Story / Attribute": _story_attribute(candidate),
                "Latest Date": _friendly_value(candidate.get("latest_date")),
                "Current Value": _friendly_value(candidate.get("current_value")),
                "Z-Score": _friendly_value(candidate.get("z_score")),
                "Mean": _friendly_value(candidate.get("mean")),
                "Low": _friendly_value(candidate.get("low")),
                "High": _friendly_value(candidate.get("high")),
                "Observations": _friendly_value(candidate.get("observations")),
                "View in Workbook": "View embedded chart" if chart_path else "",
                "Chart File": _relative_chart_path(chart_path, package_dir),
            }
        )
    return rows


def _summary_sheet_frame(summary: pd.DataFrame, side: str) -> pd.DataFrame:
    frame = summary[summary["Side"] == side.upper()].copy()
    if "Side" in frame.columns:
        frame = frame.drop(columns=["Side"])
    return frame.reset_index(drop=True)


def _build_run_summary_frame(
    package_dir: Path,
    buy_candidates: pd.DataFrame,
    sell_candidates: pd.DataFrame,
    chart_metadata: pd.DataFrame,
    *,
    generated_at: str,
    top_n: int | None,
    include_full_history: bool,
    value_label: str,
    ranking_diagnostics: dict[str, Any],
) -> pd.DataFrame:
    charts_generated = int((chart_metadata["status"] == "ok").sum()) if not chart_metadata.empty else 0
    date_min = ranking_diagnostics.get("date_min")
    date_max = ranking_diagnostics.get("date_max")
    rows = [
        ("Report Generated", generated_at),
        ("Output Folder", str(package_dir)),
        ("Recommended First File", "CLIENT_REPORT_BUY_SELL.xlsx"),
        ("BUY Ideas", int(len(buy_candidates))),
        ("SELL Ideas", int(len(sell_candidates))),
        ("Charts Generated", charts_generated),
        ("Charts Requested", int(len(chart_metadata))),
        ("Top N Requested", "" if top_n is None else int(top_n)),
        ("Ranking Basis", ranking_diagnostics.get("ranking_basis", "")),
        ("Value Label", value_label),
        (
            "Active Universe Focus",
            "2019+ origination years"
            if ranking_diagnostics.get("use_default_focus")
            else "Full history / manual filters",
        ),
        ("Scored Story Slices", ranking_diagnostics.get("scored_slices", "")),
        ("Filtered Rows", ranking_diagnostics.get("filtered_rows", "")),
        ("Minimum Observations", ranking_diagnostics.get("min_observations", "")),
        ("Date Range", f"{date_min or ''} to {date_max or ''}".strip()),
        ("Chart History", "Full chart history" if include_full_history else "Active universe only"),
        ("Next Step", "Review Buy Ideas and Sell Ideas sheets, then open linked chart files."),
    ]
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def _apply_sheet_formatting(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if worksheet.max_row >= 1:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

    integer_headers = {"Rank", "Vintage", "Observations"}
    decimal_headers = {"Current Value", "Z-Score", "Mean", "Low", "High"}
    text_headers = {"Coupon"}
    wrap_headers = {"Story Name", "Story / Attribute", "View in Workbook", "Chart File", "Value"}
    headers = [worksheet.cell(row=1, column=col_idx).value for col_idx in range(1, worksheet.max_column + 1)]

    for col_idx, header in enumerate(headers, start=1):
        header_text = str(header or "")
        max_length = len(header_text)
        for row_idx in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
            if header_text in integer_headers:
                cell.number_format = "0"
            elif header_text in decimal_headers:
                cell.number_format = "0.000"
            elif header_text in text_headers:
                cell.number_format = "@"
            elif header_text == "Latest Date":
                cell.number_format = "yyyy-mm-dd"
            cell.alignment = Alignment(vertical="top", wrap_text=header_text in wrap_headers)

        worksheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_length + 2, 11), 42)


def _apply_chart_hyperlinks(
    worksheet,
    *,
    chart_column_name: str = "Chart File",
    embedded_column_name: str = "View in Workbook",
) -> None:
    """Link story/chart cells to PNGs and embedded workbook chart sheets."""

    if worksheet.max_row < 2:
        return

    headers = {
        str(worksheet.cell(row=1, column=col_idx).value or ""): col_idx
        for col_idx in range(1, worksheet.max_column + 1)
    }
    story_col = headers.get("Story Name")
    chart_col = headers.get(chart_column_name)
    embedded_col = headers.get(embedded_column_name)
    if not chart_col:
        return

    hyperlink_font = Font(color="0563C1", underline="single")
    for row_idx in range(2, worksheet.max_row + 1):
        chart_cell = worksheet.cell(row=row_idx, column=chart_col)
        chart_path = str(chart_cell.value or "").strip()
        if not chart_path:
            continue
        chart_target = chart_path.replace("\\", "/")
        chart_cell.hyperlink = chart_target
        chart_cell.style = "Hyperlink"
        chart_cell.font = hyperlink_font
        if story_col:
            story_cell = worksheet.cell(row=row_idx, column=story_col)
            if story_cell.value not in (None, ""):
                story_cell.hyperlink = chart_target
                story_cell.style = "Hyperlink"
                story_cell.font = hyperlink_font
        if embedded_col:
            embedded_cell = worksheet.cell(row=row_idx, column=embedded_col)
            if embedded_cell.value not in (None, ""):
                rank_col = headers.get("Rank")
                if not rank_col:
                    continue
                try:
                    rank = int(worksheet.cell(row=row_idx, column=rank_col).value)
                except (TypeError, ValueError):
                    continue
                side = str(worksheet.title).split()[0].lower()
                embedded_cell.hyperlink = _embedded_chart_target(side, rank)
                embedded_cell.style = "Hyperlink"
                embedded_cell.font = hyperlink_font


def _embedded_chart_sheet_name(side: str) -> str:
    return "Buy Charts" if str(side).lower() == "buy" else "Sell Charts"


def _embedded_chart_anchor_row(rank: int) -> int:
    return 1 + (max(int(rank), 1) - 1) * 30


def _embedded_chart_target(side: str, rank: int) -> str:
    sheet_name = _embedded_chart_sheet_name(side)
    return f"#'{sheet_name}'!A{_embedded_chart_anchor_row(rank)}"


def _resolve_chart_path(chart_path: Any, package_dir: Path) -> Path | None:
    if chart_path in (None, "") or pd.isna(chart_path):
        return None
    path = Path(str(chart_path))
    if not path.is_absolute():
        path = package_dir / path
    return path if path.exists() else None


def _add_embedded_chart_sheets(workbook, chart_metadata: pd.DataFrame, package_dir: Path) -> None:
    """Embed chart PNGs so the workbook still works when shared by itself."""

    if chart_metadata.empty:
        return

    title_fill = PatternFill("solid", fgColor="EAF2F8")
    title_font = Font(bold=True, color="1F4E78", size=12)

    for side in ("buy", "sell"):
        if "side" not in chart_metadata.columns or "status" not in chart_metadata.columns:
            continue
        rows = chart_metadata[
            (chart_metadata["side"].astype(str).str.lower() == side)
            & (chart_metadata["status"].astype(str) == "ok")
        ].copy()
        if rows.empty:
            continue

        sheet_name = _embedded_chart_sheet_name(side)
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.sheet_view.showGridLines = False
        worksheet.column_dimensions["A"].width = 24
        worksheet.column_dimensions["B"].width = 96
        worksheet.column_dimensions["C"].width = 18

        for _, row in rows.sort_values("rank").iterrows():
            try:
                rank = int(row.get("rank"))
            except (TypeError, ValueError):
                continue
            anchor_row = _embedded_chart_anchor_row(rank)
            chart_path = _resolve_chart_path(row.get("chart_path"), package_dir)
            title = str(row.get("raw_story_name") or f"{side.upper()} idea #{rank}")

            worksheet.cell(row=anchor_row, column=1, value=f"{side.upper()} #{rank}")
            worksheet.cell(row=anchor_row, column=2, value=title)
            worksheet.cell(row=anchor_row + 1, column=1, value="PNG file")
            worksheet.cell(
                row=anchor_row + 1,
                column=2,
                value=_relative_chart_path(row.get("chart_path"), package_dir),
            )
            for cell in worksheet[anchor_row]:
                cell.fill = title_fill
                cell.font = title_font
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            worksheet.row_dimensions[anchor_row].height = 24

            if chart_path:
                try:
                    image = ExcelImage(str(chart_path))
                    image.width = 880
                    image.height = 440
                    worksheet.add_image(image, f"A{anchor_row + 3}")
                except Exception:
                    worksheet.cell(row=anchor_row + 3, column=1, value="Chart image could not be embedded.")


def _format_report_workbook(
    workbook,
    *,
    chart_metadata: pd.DataFrame | None = None,
    package_dir: Path | None = None,
) -> None:
    if chart_metadata is not None and package_dir is not None:
        _add_embedded_chart_sheets(workbook, chart_metadata, package_dir)

    for sheet_name in workbook.sheetnames:
        if sheet_name in {"Buy Charts", "Sell Charts"}:
            continue
        worksheet = workbook[sheet_name]
        _apply_sheet_formatting(worksheet)

    if "Run Summary" in workbook.sheetnames:
        worksheet = workbook["Run Summary"]
        worksheet.column_dimensions["A"].width = 28
        worksheet.column_dimensions["B"].width = 72

    for sheet_name in ("Buy Ideas", "Sell Ideas"):
        if sheet_name in workbook.sheetnames:
            _apply_chart_hyperlinks(workbook[sheet_name])

    if "Technical Metadata" in workbook.sheetnames:
        _apply_chart_hyperlinks(workbook["Technical Metadata"], chart_column_name="chart_path")


def _write_excel_friendly_csv(frame: pd.DataFrame, path: Path) -> None:
    """Write a CSV that opens into columns when double-clicked in Excel."""

    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        # Excel honors this delimiter hint even when Windows locale expects semicolons.
        handle.write("sep=,\n")
        frame.to_csv(handle, index=False)


def _build_readable_exports(
    buy_candidates: pd.DataFrame,
    sell_candidates: pd.DataFrame,
    chart_metadata: pd.DataFrame,
    package_dir: Path,
    *,
    generated_at: str,
    top_n: int | None,
    include_full_history: bool,
    value_label: str,
    ranking_diagnostics: dict[str, Any],
) -> tuple[Path, Path | None, str | None]:
    """Write client-readable BUY/SELL summary files next to technical outputs."""

    chart_lookup: dict[tuple[str, int], pd.Series] = {}
    if not chart_metadata.empty:
        for _, row in chart_metadata.iterrows():
            try:
                key = (str(row.get("side", "")).lower(), int(row.get("rank")))
            except (TypeError, ValueError):
                continue
            chart_lookup[key] = row

    summary_rows = []
    summary_rows.extend(_candidate_summary_rows(buy_candidates, chart_lookup, package_dir, side="buy"))
    summary_rows.extend(_candidate_summary_rows(sell_candidates, chart_lookup, package_dir, side="sell"))

    summary_columns = [
        "Side",
        "Rank",
        "Story Name",
        "Program",
        "Vintage",
        "Coupon",
        "Story / Attribute",
        "Latest Date",
        "Current Value",
        "Z-Score",
        "Mean",
        "Low",
        "High",
        "Observations",
        "View in Workbook",
        "Chart File",
    ]
    summary = pd.DataFrame(summary_rows, columns=summary_columns)
    summary_path = package_dir / "diagnostics" / "technical_buy_sell_summary.csv"
    _write_excel_friendly_csv(summary, summary_path)
    buy_summary = _summary_sheet_frame(summary, "BUY")
    sell_summary = _summary_sheet_frame(summary, "SELL")
    run_summary = _build_run_summary_frame(
        package_dir,
        buy_candidates,
        sell_candidates,
        chart_metadata,
        generated_at=generated_at,
        top_n=top_n,
        include_full_history=include_full_history,
        value_label=value_label,
        ranking_diagnostics=ranking_diagnostics,
    )

    technical_buy = buy_candidates.copy()
    technical_buy.insert(0, "side", "buy")
    technical_buy.insert(1, "rank", range(1, len(technical_buy) + 1))
    technical_sell = sell_candidates.copy()
    technical_sell.insert(0, "side", "sell")
    technical_sell.insert(1, "rank", range(1, len(technical_sell) + 1))
    technical = pd.concat([technical_buy, technical_sell], ignore_index=True)
    if not chart_metadata.empty:
        chart_slim = chart_metadata[["side", "rank", "status", "chart_path"]].copy()
        chart_slim["chart_path"] = chart_slim["chart_path"].apply(lambda value: _relative_chart_path(value, package_dir))
        technical = technical.merge(chart_slim, on=["side", "rank"], how="left")

    workbook_path = package_dir / "CLIENT_REPORT_BUY_SELL.xlsx"
    workbook_error = None
    try:
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            run_summary.to_excel(writer, sheet_name="Run Summary", index=False)
            buy_summary.to_excel(writer, sheet_name="Buy Ideas", index=False)
            sell_summary.to_excel(writer, sheet_name="Sell Ideas", index=False)
            technical.to_excel(writer, sheet_name="Technical Metadata", index=False)
            _format_report_workbook(writer.book, chart_metadata=chart_metadata, package_dir=package_dir)
    except Exception as exc:
        workbook_error = str(exc)
        workbook_path = None

    return summary_path, workbook_path, workbook_error


def build_idea_chart_package(
    normalized_frame: pd.DataFrame,
    ranking_result: RankingResult,
    output_dir: str | Path,
    *,
    top_n: int | None = None,
    include_full_history: bool = True,
    value_label: str = "Value",
) -> IdeaChartPackageResult:
    """Create CSVs, charts, and metadata for ranked BUY/SELL ideas.

    Ranking is intentionally not recalculated here. The provided `RankingResult`
    owns candidate selection, including active-universe defaults.
    """

    package_dir = ensure_dir(output_dir)
    ensure_dir(package_dir / "charts" / "buy")
    ensure_dir(package_dir / "charts" / "sell")
    diagnostics_dir = ensure_dir(package_dir / "diagnostics")

    normalized = add_comparable_keys(normalized_frame)
    normalized["date"] = pd.to_datetime(normalized["date"], errors="coerce")
    normalized["value"] = pd.to_numeric(normalized["value"], errors="coerce")
    normalized = normalized.dropna(subset=["date", "value"])

    buy_candidates = _candidate_subset(ranking_result.buy_candidates, top_n)
    sell_candidates = _candidate_subset(ranking_result.sell_candidates, top_n)

    buy_candidates_path = diagnostics_dir / "technical_buy_candidates.csv"
    sell_candidates_path = diagnostics_dir / "technical_sell_candidates.csv"
    buy_candidates.to_csv(buy_candidates_path, index=False)
    sell_candidates.to_csv(sell_candidates_path, index=False)

    chart_records = []
    chart_records.extend(
        _build_chart_records(
            normalized,
            buy_candidates,
            package_dir,
            side="buy",
            include_full_history=include_full_history,
            value_label=value_label,
        )
    )
    chart_records.extend(
        _build_chart_records(
            normalized,
            sell_candidates,
            package_dir,
            side="sell",
            include_full_history=include_full_history,
            value_label=value_label,
        )
    )

    chart_metadata = pd.DataFrame(chart_records)
    chart_metadata_path = diagnostics_dir / "chart_metadata.csv"
    chart_metadata.to_csv(chart_metadata_path, index=False)
    generated_at = now_iso()
    readable_summary_path, readable_workbook_path, readable_workbook_error = _build_readable_exports(
        buy_candidates,
        sell_candidates,
        chart_metadata,
        package_dir,
        generated_at=generated_at,
        top_n=top_n,
        include_full_history=include_full_history,
        value_label=value_label,
        ranking_diagnostics=ranking_result.diagnostics,
    )

    ok_charts = int((chart_metadata["status"] == "ok").sum()) if not chart_metadata.empty else 0
    metadata = {
        "generated_at": generated_at,
        "output_dir": str(package_dir),
        "top_n": None if top_n is None else int(top_n),
        "include_full_history": bool(include_full_history),
        "value_label": value_label,
        "buy_candidates": int(len(buy_candidates)),
        "sell_candidates": int(len(sell_candidates)),
        "charts_generated": ok_charts,
        "charts_requested": int(len(chart_metadata)),
        "client_readable_outputs": {
            "client_report_xlsx": str(readable_workbook_path) if readable_workbook_path else None,
            "technical_summary_csv": str(readable_summary_path),
            "excel_error": readable_workbook_error,
        },
        "ranking_diagnostics": ranking_result.diagnostics,
        "paths": {
            "technical_buy_candidates": str(buy_candidates_path),
            "technical_sell_candidates": str(sell_candidates_path),
            "technical_buy_sell_summary": str(readable_summary_path),
            "client_report_workbook": str(readable_workbook_path) if readable_workbook_path else None,
            "chart_metadata": str(chart_metadata_path),
        },
    }
    package_metadata_path = diagnostics_dir / "package_metadata.json"
    write_json(package_metadata_path, metadata)

    return IdeaChartPackageResult(
        output_dir=package_dir,
        buy_candidates_path=buy_candidates_path,
        sell_candidates_path=sell_candidates_path,
        chart_metadata_path=chart_metadata_path,
        package_metadata_path=package_metadata_path,
        chart_metadata=chart_metadata,
        metadata=metadata,
        readable_summary_path=readable_summary_path,
        readable_workbook_path=readable_workbook_path,
    )
