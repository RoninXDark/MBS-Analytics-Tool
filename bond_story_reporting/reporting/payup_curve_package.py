from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from hashlib import sha1
import math
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import pandas as pd
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..utils import ensure_dir, now_iso, slugify, write_json


SUPPORTED_PAYUP_CHART_TYPES = {"time_series", "latest_snapshot_curve"}


@dataclass(frozen=True)
class PayupCurvePackageResult:
    output_dir: Path
    index_workbook_path: Path | None
    chart_metadata_path: Path
    package_metadata_path: Path
    parse_diagnostics_path: Path
    chart_metadata: pd.DataFrame
    metadata: dict[str, Any]


def build_payup_curve_package(
    normalized_frame: pd.DataFrame,
    output_dir: str | Path,
    *,
    report_name: str,
    chart_requests: list[dict[str, Any]],
    parse_diagnostics: dict[str, Any] | None = None,
    input_path: str | Path | None = None,
    max_charts_per_request: int = 50,
    max_total_charts: int = 250,
    max_series_per_chart: int = 50,
) -> PayupCurvePackageResult:
    """Generate a client-readable canned payup curve package.

    Phase 7B intentionally keeps this package independent from Tkinter and from
    BUY/SELL ranking. It consumes the normalized dataset and saved chart
    requests, then writes charts plus an index workbook.
    """

    package_dir = ensure_dir(output_dir)
    chart_dir = ensure_dir(package_dir / "charts")
    diagnostics_dir = ensure_dir(package_dir / "diagnostics")

    normalized = _prepare_normalized_frame(normalized_frame)
    requests = [_normalize_chart_request(item, idx + 1) for idx, item in enumerate(chart_requests or [])]

    records: list[dict[str, Any]] = []
    request_summaries: list[dict[str, Any]] = []
    total_charts = 0

    for request in requests:
        if total_charts >= max_total_charts:
            request_summaries.append(
                _request_summary(
                    request,
                    status="skipped_total_chart_limit",
                    row_count=0,
                    series_count=0,
                    charts_generated=0,
                    warning=f"Skipped because max_total_charts={max_total_charts} was reached.",
                )
            )
            continue

        if request["chart_type"] not in SUPPORTED_PAYUP_CHART_TYPES:
            request_summaries.append(
                _request_summary(
                    request,
                    status="unsupported_chart_type",
                    row_count=0,
                    series_count=0,
                    charts_generated=0,
                    warning=(
                        f"Unsupported chart_type '{request['chart_type']}'. "
                        f"Supported types: {', '.join(sorted(SUPPORTED_PAYUP_CHART_TYPES))}."
                    ),
                )
            )
            continue

        filtered = _apply_request_filters(normalized, request)
        if filtered.empty:
            request_summaries.append(
                _request_summary(
                    request,
                    status="no_matching_data",
                    row_count=0,
                    series_count=0,
                    charts_generated=0,
                    warning="No normalized rows matched this request.",
                )
            )
            continue

        charts_allowed = min(max_charts_per_request, max_total_charts - total_charts)
        if charts_allowed <= 0:
            request_summaries.append(
                _request_summary(
                    request,
                    status="skipped_total_chart_limit",
                    row_count=int(len(filtered)),
                    series_count=int(filtered["raw_story_name"].nunique()),
                    charts_generated=0,
                    warning=f"Skipped because max_total_charts={max_total_charts} was reached.",
                )
            )
            continue

        if request["chart_type"] == "time_series":
            record = _generate_time_series_chart(
                filtered,
                chart_dir,
                request,
                package_dir=package_dir,
                max_series_per_chart=max_series_per_chart,
            )
        else:
            record = _generate_latest_snapshot_curve(
                filtered,
                chart_dir,
                request,
                package_dir=package_dir,
            )

        records.append(record)
        total_charts += 1 if record["status"] in {"ok", "ok_truncated_series"} else 0
        request_summaries.append(
            _request_summary(
                request,
                status=record["status"],
                row_count=int(len(filtered)),
                series_count=int(record.get("series_count") or 0),
                charts_generated=1 if record["status"] in {"ok", "ok_truncated_series"} else 0,
                warning=record.get("warning"),
            )
        )

    chart_metadata = pd.DataFrame(records)
    chart_metadata_path = diagnostics_dir / "chart_metadata.csv"
    chart_metadata.to_csv(chart_metadata_path, index=False)

    parse_diagnostics_path = diagnostics_dir / "parse_diagnostics.json"
    write_json(parse_diagnostics_path, parse_diagnostics or {})

    generated_at = now_iso()
    charts_generated = int(
        chart_metadata["status"].isin(["ok", "ok_truncated_series"]).sum()
    ) if not chart_metadata.empty else 0
    skipped_requests = int(
        sum(1 for item in request_summaries if int(item.get("Charts Generated") or 0) == 0)
    )
    warnings = [str(item["Warning"]) for item in request_summaries if str(item.get("Warning") or "").strip()]
    status = "completed"
    if charts_generated == 0:
        status = "completed_no_charts"
    elif warnings or skipped_requests:
        status = "completed_with_warnings"

    index_workbook_path, workbook_error = _write_index_workbook(
        package_dir,
        report_name=report_name,
        generated_at=generated_at,
        input_path=input_path,
        request_summaries=request_summaries,
        chart_metadata=chart_metadata,
        parse_diagnostics=parse_diagnostics or {},
        charts_generated=charts_generated,
        skipped_requests=skipped_requests,
        warnings=warnings,
    )

    metadata = {
        "generated_at": generated_at,
        "status": status,
        "input_file": str(input_path) if input_path is not None else None,
        "report_name": report_name,
        "chart_request_count": int(len(requests)),
        "charts_generated": charts_generated,
        "skipped_requests": skipped_requests,
        "output_folder": str(package_dir),
        "max_charts_per_request": int(max_charts_per_request),
        "max_total_charts": int(max_total_charts),
        "max_series_per_chart": int(max_series_per_chart),
        "warnings": warnings,
        "paths": {
            "client_report_index": str(index_workbook_path) if index_workbook_path else None,
            "chart_metadata": str(chart_metadata_path),
            "parse_diagnostics": str(parse_diagnostics_path),
            "package_metadata": str(diagnostics_dir / "package_metadata.json"),
        },
        "excel_error": workbook_error,
    }
    package_metadata_path = diagnostics_dir / "package_metadata.json"
    write_json(package_metadata_path, metadata)

    return PayupCurvePackageResult(
        output_dir=package_dir,
        index_workbook_path=index_workbook_path,
        chart_metadata_path=chart_metadata_path,
        package_metadata_path=package_metadata_path,
        parse_diagnostics_path=parse_diagnostics_path,
        chart_metadata=chart_metadata,
        metadata=metadata,
    )


def _prepare_normalized_frame(frame: pd.DataFrame) -> pd.DataFrame:
    work = frame.copy()
    work["date"] = pd.to_datetime(work["date"], errors="coerce")
    work["value"] = pd.to_numeric(work["value"], errors="coerce")
    work["coupon"] = pd.to_numeric(work.get("coupon"), errors="coerce")
    work["origination_year"] = pd.to_numeric(work.get("origination_year"), errors="coerce")
    return work.dropna(subset=["date", "value"]).reset_index(drop=True)


def _normalize_chart_request(raw: dict[str, Any], request_index: int) -> dict[str, Any]:
    request = dict(raw or {})
    request["request_index"] = request_index
    request["name"] = str(request.get("name") or f"Chart Request {request_index}").strip()
    request["chart_type"] = str(request.get("chart_type") or "time_series").strip().lower()
    return request


def _apply_request_filters(frame: pd.DataFrame, request: dict[str, Any]) -> pd.DataFrame:
    work = frame.copy()
    filters = {
        "program": _as_str_list(request.get("programs") or request.get("program")),
        "geography": _as_str_list(request.get("geographies") or request.get("geo")),
        "loan_balance_bucket": _as_str_list(
            request.get("loan_balance_buckets") or request.get("loan_balances")
        ),
        "story_family": _as_str_list(request.get("story_families") or request.get("story_family")),
        "story_tag": _as_str_list(request.get("story_tags") or request.get("story_tag")),
        "raw_story_name": _as_str_list(request.get("raw_story_names") or request.get("stories")),
    }
    for column, values in filters.items():
        if values and column in work.columns:
            work = _filter_str_column(work, column, values)

    coupon_value = request.get("coupons") if "coupons" in request else request.get("coupon")
    if not _is_all_token(coupon_value):
        coupon_values = _as_float_list(coupon_value)
        if coupon_values:
            work = work[work["coupon"].isin(coupon_values)].copy()

    coupon_min = _as_optional_float(request.get("coupon_min"))
    coupon_max = _as_optional_float(request.get("coupon_max"))
    if coupon_min is not None:
        work = work[work["coupon"] >= coupon_min].copy()
    if coupon_max is not None:
        work = work[work["coupon"] <= coupon_max].copy()

    years = _as_int_list(request.get("origination_years") or request.get("years"))
    if years:
        work = work[work["origination_year"].isin(years)].copy()

    year_min = _as_optional_int(request.get("origination_year_min") or request.get("year_min"))
    year_max = _as_optional_int(request.get("origination_year_max") or request.get("year_max"))
    if year_min is not None:
        work = work[work["origination_year"] >= year_min].copy()
    if year_max is not None:
        work = work[work["origination_year"] <= year_max].copy()

    return work.sort_values(["date", "raw_story_name"]).reset_index(drop=True)


def _filter_str_column(frame: pd.DataFrame, column: str, values: list[str]) -> pd.DataFrame:
    lookup = {value.strip().upper() for value in values if value.strip()}
    if not lookup:
        return frame
    series = frame[column].fillna("").astype(str).str.strip().str.upper()
    return frame[series.isin(lookup)].copy()


def _generate_time_series_chart(
    frame: pd.DataFrame,
    chart_dir: Path,
    request: dict[str, Any],
    *,
    package_dir: Path,
    max_series_per_chart: int,
) -> dict[str, Any]:
    groups = []
    for story_name, series in frame.groupby("raw_story_name", dropna=False):
        clean = series.dropna(subset=["date", "value"]).sort_values("date")
        if clean.empty:
            continue
        groups.append((str(story_name), clean, int(len(clean))))

    if not groups:
        return _chart_record(request, status="no_series_data", package_dir=package_dir)

    groups.sort(key=lambda item: (-item[2], item[0].lower()))
    warning = None
    total_series = len(groups)
    if total_series > max_series_per_chart:
        groups = groups[:max_series_per_chart]
        warning = f"Series truncated from {total_series} to max_series_per_chart={max_series_per_chart}."

    filename = _chart_filename(request, suffix="time_series")
    chart_path = chart_dir / filename

    fig, ax = plt.subplots(figsize=(13, 7))
    colors = plt.cm.tab20.colors
    for idx, (story_name, series, _) in enumerate(groups):
        plot_series = series.groupby("date", as_index=False)["value"].mean().sort_values("date")
        label = _short_label(_series_label(story_name, series), max_length=52)
        ax.plot(
            plot_series["date"],
            plot_series["value"],
            linewidth=1.8,
            marker="o",
            markersize=2.5,
            color=colors[idx % len(colors)],
            label=label,
        )

    chart_title = _chart_title(request, frame)
    ax.set_title(chart_title, fontsize=13, fontweight="bold")
    ax.set_xlabel("Date")
    ax.set_ylabel("Payup / Current Value")
    ax.grid(alpha=0.25)
    if len(groups) <= 24:
        ax.legend(loc="best", fontsize=7, ncol=1 if len(groups) <= 12 else 2)
    else:
        ax.text(
            0.01,
            0.01,
            f"{len(groups)} series plotted; legend hidden for readability.",
            transform=ax.transAxes,
            fontsize=8,
            color="#555555",
        )
    fig.autofmt_xdate()
    fig.tight_layout()
    fig.savefig(chart_path, dpi=180, bbox_inches="tight")
    plt.close(fig)

    status = "ok_truncated_series" if warning else "ok"
    return _chart_record(
        request,
        status=status,
        package_dir=package_dir,
        chart_path=chart_path,
        row_count=int(len(frame)),
        series_count=len(groups),
        matched_series_count=total_series,
        latest_date=_safe_date(frame["date"].max()),
        chart_title=chart_title,
        warning=warning,
    )


def _generate_latest_snapshot_curve(
    frame: pd.DataFrame,
    chart_dir: Path,
    request: dict[str, Any],
    *,
    package_dir: Path,
) -> dict[str, Any]:
    clean = frame.dropna(subset=["coupon", "date", "value"]).copy()
    if clean.empty:
        return _chart_record(
            request,
            status="no_coupon_data",
            package_dir=package_dir,
            warning="No rows with coupon/date/value were available for snapshot curve.",
        )

    latest_rows = (
        clean.sort_values("date")
        .groupby("raw_story_name", as_index=False)
        .tail(1)
        .reset_index(drop=True)
    )
    x_axis, axis_warning = _snapshot_axis_mode(request, latest_rows)
    if axis_warning:
        return _chart_record(
            request,
            status="ambiguous_snapshot_axes",
            package_dir=package_dir,
            row_count=int(len(frame)),
            series_count=int(latest_rows["raw_story_name"].nunique()),
            matched_series_count=int(latest_rows["raw_story_name"].nunique()),
            warning=axis_warning,
        )

    if x_axis == "vintage":
        curve = (
            latest_rows.dropna(subset=["origination_year"])
            .groupby("origination_year", as_index=False)
            .agg(value=("value", "mean"), story_count=("raw_story_name", "nunique"), latest_date=("date", "max"))
            .sort_values("origination_year")
        )
        x_column = "origination_year"
        x_label = "Vintage / Origination Year"
        suffix = "vintage_snapshot"
    else:
        curve = (
            latest_rows.groupby("coupon", as_index=False)
            .agg(value=("value", "mean"), story_count=("raw_story_name", "nunique"), latest_date=("date", "max"))
            .sort_values("coupon")
        )
        x_column = "coupon"
        x_label = "Coupon"
        suffix = "snapshot_curve"

    if curve.empty:
        return _chart_record(request, status="no_coupon_data", package_dir=package_dir)

    filename = _chart_filename(request, suffix=suffix)
    chart_path = chart_dir / filename

    fig, ax = plt.subplots(figsize=(11, 6))
    ax.plot(curve[x_column], curve["value"], color="#1f6feb", linewidth=2.2, marker="o", markersize=6)
    for _, row in curve.iterrows():
        ax.annotate(
            f"{float(row['value']):.3f}",
            (float(row[x_column]), float(row["value"])),
            textcoords="offset points",
            xytext=(0, 7),
            ha="center",
            fontsize=8,
        )
    chart_title = _chart_title(request, frame)
    ax.set_title(chart_title, fontsize=13, fontweight="bold")
    ax.set_xlabel(x_label)
    ax.set_ylabel("Latest Payup / Current Value")
    if x_axis == "vintage":
        ticks = [int(value) for value in curve[x_column].dropna().tolist()]
        if ticks:
            ax.set_xticks(ticks)
            ax.set_xticklabels([str(tick) for tick in ticks])
            ax.set_xlim(min(ticks) - 0.35, max(ticks) + 0.35)
    else:
        ticks = _coupon_axis_ticks(curve[x_column], request)
        if ticks:
            ax.set_xticks(ticks)
            ax.set_xticklabels([_format_coupon_tick(tick) for tick in ticks])
            ax.set_xlim(min(ticks) - 0.08, max(ticks) + 0.08)
    ax.grid(alpha=0.25)
    fig.tight_layout()
    fig.savefig(chart_path, dpi=180, bbox_inches="tight")
    plt.close(fig)

    return _chart_record(
        request,
        status="ok",
        package_dir=package_dir,
        chart_path=chart_path,
        row_count=int(len(frame)),
        series_count=int(latest_rows["raw_story_name"].nunique()),
        matched_series_count=int(latest_rows["raw_story_name"].nunique()),
        latest_date=_safe_date(curve["latest_date"].max()),
        coupon_count=int(len(curve)),
        x_axis=x_axis,
        chart_title=chart_title,
    )


def _chart_record(
    request: dict[str, Any],
    *,
    status: str,
    package_dir: Path,
    chart_path: Path | None = None,
    row_count: int = 0,
    series_count: int = 0,
    matched_series_count: int = 0,
    latest_date: str | None = None,
    coupon_count: int | None = None,
    x_axis: str | None = None,
    chart_title: str | None = None,
    warning: str | None = None,
) -> dict[str, Any]:
    return {
        "request_index": int(request["request_index"]),
        "request_name": request["name"],
        "chart_type": request["chart_type"],
        "chart_title": chart_title or _chart_title(request),
        "status": status,
        "chart_path": str(chart_path) if chart_path else None,
        "relative_chart_path": _relative_path(chart_path, package_dir) if chart_path else "",
        "story_attribute": _request_story_attribute(request),
        "row_count": int(row_count),
        "series_count": int(series_count),
        "matched_series_count": int(matched_series_count),
        "coupon_count": coupon_count,
        "x_axis": x_axis,
        "latest_date": latest_date,
        "warning": warning,
        "filters": _request_filter_summary(request),
    }


def _request_summary(
    request: dict[str, Any],
    *,
    status: str,
    row_count: int,
    series_count: int,
    charts_generated: int,
    warning: str | None = None,
) -> dict[str, Any]:
    return {
        "Request #": int(request["request_index"]),
        "Request Name": request["name"],
        "Chart Type": request["chart_type"],
        "Story / Attribute": _request_story_attribute(request),
        "X-Axis": request.get("x_axis") or request.get("snapshot_x_axis") or "",
        "Status": status,
        "Rows Matched": int(row_count),
        "Series Matched": int(series_count),
        "Charts Generated": int(charts_generated),
        "Warning": warning or "",
        "Filters": _request_filter_summary(request),
    }


def _request_filter_summary(request: dict[str, Any]) -> str:
    keys = [
        "programs",
        "coupons",
        "coupon_min",
        "coupon_max",
        "x_axis",
        "snapshot_x_axis",
        "origination_years",
        "origination_year_min",
        "origination_year_max",
        "geographies",
        "loan_balance_buckets",
        "story_families",
        "story_tags",
        "raw_story_names",
    ]
    parts = []
    for key in keys:
        value = request.get(key)
        if value not in (None, "", [], {}):
            parts.append(f"{key}={value}")
    return "; ".join(parts)


def _request_story_attribute(request: dict[str, Any]) -> str:
    keys = [
        "geographies",
        "loan_balance_buckets",
        "story_tags",
        "story_families",
        "raw_story_names",
    ]
    parts: list[str] = []
    for key in keys:
        value = request.get(key)
        if value in (None, "", [], {}):
            continue
        for item in _as_str_list(value):
            if item.lower() in {"all", "*", "any"}:
                continue
            if item and item not in parts:
                parts.append(item)
    return " / ".join(parts) if parts else "All stories"


def _snapshot_axis_mode(request: dict[str, Any], latest_rows: pd.DataFrame) -> tuple[str, str | None]:
    """Choose a safe X-axis for snapshot charts without silently mixing dimensions."""

    explicit = str(request.get("x_axis") or request.get("snapshot_x_axis") or "").strip().lower()
    if explicit in {"origination_year", "origination_years", "year", "years", "vintage", "vintages"}:
        preferred = "vintage"
    elif explicit in {"coupon", "coupons"}:
        preferred = "coupon"
    else:
        preferred = ""

    coupon_count = _unique_numeric_count(latest_rows, "coupon")
    vintage_count = _unique_numeric_count(latest_rows, "origination_year")

    if preferred == "coupon":
        if vintage_count > 1:
            return (
                "coupon",
                "Coupon snapshot has more than one vintage. Add one origination_year, "
                "or use one coupon with multiple years for a vintage comparison.",
            )
        return "coupon", None

    if preferred == "vintage":
        if coupon_count > 1:
            return (
                "vintage",
                "Vintage snapshot has more than one coupon. Add one coupon, "
                "or use one origination_year with multiple coupons for a coupon curve.",
            )
        return "vintage", None

    if vintage_count > 1 and coupon_count <= 1:
        return "vintage", None
    if vintage_count <= 1:
        return "coupon", None
    return (
        "coupon",
        "Snapshot request matched multiple coupons and multiple vintages. "
        "Specify a single vintage for a coupon curve, or a single coupon for a vintage curve.",
    )


def _unique_numeric_count(frame: pd.DataFrame, column: str) -> int:
    if column not in frame.columns:
        return 0
    return int(pd.to_numeric(frame[column], errors="coerce").dropna().nunique())


def _series_label(story_name: str, series: pd.DataFrame) -> str:
    clean = series.sort_values("date")
    row = clean.iloc[-1] if not clean.empty else {}
    parts: list[str] = []
    program = _safe_cell(row, "program")
    coupon = _safe_numeric_cell(row, "coupon")
    year = _safe_numeric_cell(row, "origination_year")
    if program:
        parts.append(program)
    if coupon is not None:
        parts.append(_format_coupon_tick(coupon))
    if year is not None:
        parts.append(str(int(year)))

    attrs = []
    for column in ("geography", "loan_balance_bucket", "story_tag", "story_family"):
        value = _safe_cell(row, column)
        if value and value not in attrs and value not in parts:
            attrs.append(value)
    parts.extend(attrs[:3])
    return " ".join(parts) if parts else str(story_name)


def _safe_cell(row: Any, column: str) -> str:
    try:
        value = row.get(column)
    except AttributeError:
        return ""
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def _safe_numeric_cell(row: Any, column: str) -> float | None:
    try:
        value = row.get(column)
    except AttributeError:
        return None
    if value is None or pd.isna(value):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _chart_title(request: dict[str, Any], frame: pd.DataFrame | None = None) -> str:
    """Build a readable chart title from explicit filters and safe inferences."""

    base_title = str(request.get("name") or "Payup Chart").strip() or "Payup Chart"
    program_label = _request_program_label(request, frame)
    vintage_label = _request_vintage_label(request, frame)

    title = base_title
    if program_label and program_label.lower() not in title.lower():
        title = f"{program_label} {title}"
    if vintage_label and vintage_label.lower() not in title.lower():
        title = f"{title} - {vintage_label}"
    return title


def _request_program_label(request: dict[str, Any], frame: pd.DataFrame | None = None) -> str:
    programs = _as_str_list(request.get("programs") or request.get("program"))
    programs = [item for item in programs if item and item.lower() not in {"all", "*", "any"}]
    if not programs and frame is not None and "program" in frame.columns:
        values = sorted(
            {
                str(value).strip()
                for value in frame["program"].dropna().tolist()
                if str(value).strip()
            }
        )
        if len(values) == 1:
            programs = values
    if not programs:
        return ""
    if len(programs) == 1:
        return programs[0]
    return "/".join(programs[:3]) + ("+" if len(programs) > 3 else "")


def _request_vintage_label(request: dict[str, Any], frame: pd.DataFrame | None = None) -> str:
    years = _as_int_list(request.get("origination_years") or request.get("years"))
    if years:
        unique_years = sorted(set(years))
        if len(unique_years) == 1:
            return f"{unique_years[0]} Vintage"
        return f"{unique_years[0]}-{unique_years[-1]} Vintages"

    year_min = _as_optional_int(request.get("origination_year_min") or request.get("year_min"))
    year_max = _as_optional_int(request.get("origination_year_max") or request.get("year_max"))
    if year_min is not None and year_max is not None:
        return f"{year_min}-{year_max} Vintages"
    if year_min is not None:
        return f"{year_min}+ Vintages"
    if year_max is not None:
        return f"Through {year_max} Vintages"

    if frame is not None and "origination_year" in frame.columns:
        values = sorted(
            {
                int(value)
                for value in pd.to_numeric(frame["origination_year"], errors="coerce").dropna().tolist()
            }
        )
        if len(values) == 1:
            return f"{values[0]} Vintage"
    return ""


def _coupon_axis_ticks(coupons: Any, request: dict[str, Any]) -> list[float]:
    values = pd.to_numeric(pd.Series(coupons), errors="coerce").dropna()
    if values.empty:
        return []

    minimum = _as_optional_float(request.get("coupon_min"))
    maximum = _as_optional_float(request.get("coupon_max"))
    if minimum is None:
        minimum = math.floor(float(values.min()) * 2) / 2
    if maximum is None:
        maximum = math.ceil(float(values.max()) * 2) / 2
    if maximum < minimum:
        return []

    tick_count = int(round((maximum - minimum) / 0.5)) + 1
    return [round(minimum + 0.5 * idx, 2) for idx in range(tick_count)]


def _format_coupon_tick(value: float) -> str:
    return f"{float(value):.1f}".rstrip("0").rstrip(".")


def _chart_filename(request: dict[str, Any], *, suffix: str) -> str:
    key = f"{request['request_index']}|{request['name']}|{request['chart_type']}|{suffix}"
    digest = sha1(key.encode("utf-8")).hexdigest()[:8]
    clean = slugify(str(request["name"]))[:72] or "payup_curve"
    return f"{int(request['request_index']):03d}_{clean}_{suffix}_{digest}.png"


def _write_index_workbook(
    package_dir: Path,
    *,
    report_name: str,
    generated_at: str,
    input_path: str | Path | None,
    request_summaries: list[dict[str, Any]],
    chart_metadata: pd.DataFrame,
    parse_diagnostics: dict[str, Any],
    charts_generated: int,
    skipped_requests: int,
    warnings: list[str],
) -> tuple[Path | None, str | None]:
    workbook_path = package_dir / "CLIENT_REPORT_INDEX.xlsx"
    request_frame = pd.DataFrame(request_summaries)
    chart_frame = _client_chart_frame(chart_metadata, package_dir)
    technical_frame = _technical_metadata_frame(parse_diagnostics=parse_diagnostics, warnings=warnings)
    summary_frame = pd.DataFrame(
        [
            ("Report Name", report_name),
            ("Generated At", generated_at),
            ("Input File", "" if input_path is None else str(input_path)),
            ("Output Folder", str(package_dir)),
            ("Recommended First File", "CLIENT_REPORT_INDEX.xlsx"),
            ("Chart Requests", int(len(request_summaries))),
            ("Charts Generated", int(charts_generated)),
            ("Skipped / Empty Requests", int(skipped_requests)),
            ("Warnings", int(len(warnings))),
            ("Next Step", "Open Generated Charts sheet, then review PNG files in charts/."),
        ],
        columns=["Metric", "Value"],
    )

    try:
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            summary_frame.to_excel(writer, sheet_name="Run Summary", index=False)
            request_frame.to_excel(writer, sheet_name="Chart Requests Summary", index=False)
            chart_frame.to_excel(writer, sheet_name="Generated Charts", index=False)
            technical_frame.to_excel(writer, sheet_name="Technical Metadata", index=False)
            _format_workbook(writer.book, chart_metadata=chart_metadata, package_dir=package_dir)
    except Exception as exc:
        return None, str(exc)
    return workbook_path, None


def _client_chart_frame(chart_metadata: pd.DataFrame, package_dir: Path) -> pd.DataFrame:
    columns = [
        "Request #",
        "Request Name",
        "Chart Title",
        "Chart Type",
        "Story / Attribute",
        "X-Axis",
        "Status",
        "View in Workbook",
        "Chart File",
        "Latest Date",
        "Series Count",
        "Matched Series Count",
        "Coupon Count",
        "Rows Matched",
        "Warning",
    ]
    if chart_metadata.empty:
        return pd.DataFrame(columns=columns)
    rows = []
    for _, row in chart_metadata.iterrows():
        chart_path = row.get("chart_path")
        rows.append(
            {
                "Request #": row.get("request_index"),
                "Request Name": row.get("request_name"),
                "Chart Title": row.get("chart_title") or row.get("request_name"),
                "Chart Type": row.get("chart_type"),
                "Story / Attribute": row.get("story_attribute") or "All stories",
                "X-Axis": row.get("x_axis") or "",
                "Status": row.get("status"),
                "View in Workbook": "View embedded chart" if chart_path else "",
                "Chart File": _relative_path(chart_path, package_dir) if chart_path else "",
                "Latest Date": row.get("latest_date"),
                "Series Count": row.get("series_count"),
                "Matched Series Count": row.get("matched_series_count"),
                "Coupon Count": row.get("coupon_count"),
                "Rows Matched": row.get("row_count"),
                "Warning": row.get("warning"),
            }
        )
    return pd.DataFrame(rows, columns=columns)


def _technical_metadata_frame(*, parse_diagnostics: dict[str, Any], warnings: list[str]) -> pd.DataFrame:
    rows = []
    for key, value in (parse_diagnostics or {}).items():
        rows.append({"Section": "Parse Diagnostics", "Key": key, "Value": value})
    for idx, warning in enumerate(warnings, start=1):
        rows.append({"Section": "Warnings", "Key": f"warning_{idx}", "Value": warning})
    if not rows:
        rows.append({"Section": "Technical Metadata", "Key": "message", "Value": "No warnings."})
    return pd.DataFrame(rows)


def _format_workbook(
    workbook,
    *,
    chart_metadata: pd.DataFrame | None = None,
    package_dir: Path | None = None,
) -> None:
    if chart_metadata is not None and package_dir is not None:
        _add_embedded_chart_sheet(workbook, chart_metadata, package_dir)

    for sheet_name in workbook.sheetnames:
        if sheet_name == "Charts":
            continue
        worksheet = workbook[sheet_name]
        _format_sheet(worksheet)
    if "Run Summary" in workbook.sheetnames:
        workbook["Run Summary"].column_dimensions["A"].width = 28
        workbook["Run Summary"].column_dimensions["B"].width = 80
    if "Generated Charts" in workbook.sheetnames:
        _apply_chart_hyperlinks(workbook["Generated Charts"])


def _format_sheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    if worksheet.max_row >= 1:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx in range(1, worksheet.max_column + 1):
        header = str(worksheet.cell(row=1, column=col_idx).value or "")
        max_len = len(header)
        for row_idx in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=header in {"Value", "Filters", "Warning", "Story / Attribute", "Chart Title", "View in Workbook"},
            )
            if header in {"Request #", "Charts Generated", "Rows Matched", "Series Count", "Matched Series Count"}:
                cell.number_format = "0"
        worksheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 52)


def _apply_chart_hyperlinks(worksheet) -> None:
    """Make generated chart references clickable inside the canned report index."""

    if worksheet.max_row < 2:
        return

    headers = {
        str(worksheet.cell(row=1, column=col_idx).value or ""): col_idx
        for col_idx in range(1, worksheet.max_column + 1)
    }
    chart_col = headers.get("Chart File")
    name_col = headers.get("Request Name")
    embedded_col = headers.get("View in Workbook")
    if not chart_col:
        return

    hyperlink_font = Font(color="0563C1", underline="single")
    for row_idx in range(2, worksheet.max_row + 1):
        chart_cell = worksheet.cell(row=row_idx, column=chart_col)
        chart_path = str(chart_cell.value or "").strip()
        if not chart_path:
            continue
        chart_target = chart_path.replace("\\", "/")
        chart_cell.hyperlink = chart_target
        chart_cell.style = "Hyperlink"
        chart_cell.font = hyperlink_font
        if name_col:
            name_cell = worksheet.cell(row=row_idx, column=name_col)
            if name_cell.value not in (None, ""):
                name_cell.hyperlink = chart_target
                name_cell.style = "Hyperlink"
                name_cell.font = hyperlink_font
        if embedded_col:
            embedded_cell = worksheet.cell(row=row_idx, column=embedded_col)
            if embedded_cell.value not in (None, ""):
                request_col = headers.get("Request #")
                try:
                    request_index = int(worksheet.cell(row=row_idx, column=request_col).value)
                except (TypeError, ValueError):
                    continue
                embedded_cell.hyperlink = _embedded_chart_target(request_index)
                embedded_cell.style = "Hyperlink"
                embedded_cell.font = hyperlink_font


def _embedded_chart_anchor_row(index: int) -> int:
    return 1 + (max(int(index), 1) - 1) * 30


def _embedded_chart_target(index: int) -> str:
    return f"#'Charts'!A{_embedded_chart_anchor_row(index)}"


def _resolve_chart_path(path: Any, package_dir: Path) -> Path | None:
    if path in (None, "") or pd.isna(path):
        return None
    raw = Path(str(path))
    if not raw.is_absolute():
        raw = package_dir / raw
    return raw if raw.exists() else None


def _add_embedded_chart_sheet(workbook, chart_metadata: pd.DataFrame, package_dir: Path) -> None:
    """Embed canned-report chart PNGs into the workbook for one-file sharing."""

    if chart_metadata.empty or "status" not in chart_metadata.columns:
        return
    rows = chart_metadata[chart_metadata["status"].astype(str).isin(["ok", "ok_truncated_series"])].copy()
    if rows.empty:
        return

    if "Charts" in workbook.sheetnames:
        del workbook["Charts"]
    worksheet = workbook.create_sheet("Charts")
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 96
    worksheet.column_dimensions["C"].width = 18

    title_fill = PatternFill("solid", fgColor="EAF2F8")
    title_font = Font(bold=True, color="1F4E78", size=12)

    for _, row in rows.sort_values("request_index").iterrows():
        try:
            request_index = int(row.get("request_index"))
        except (TypeError, ValueError):
            continue
        anchor_row = _embedded_chart_anchor_row(request_index)
        chart_path = _resolve_chart_path(row.get("chart_path"), package_dir)
        title = str(row.get("chart_title") or row.get("request_name") or f"Chart {request_index}")

        worksheet.cell(row=anchor_row, column=1, value=f"Chart #{request_index}")
        worksheet.cell(row=anchor_row, column=2, value=title)
        worksheet.cell(row=anchor_row + 1, column=1, value="PNG file")
        worksheet.cell(row=anchor_row + 1, column=2, value=_relative_path(row.get("chart_path"), package_dir))
        for cell in worksheet[anchor_row]:
            cell.fill = title_fill
            cell.font = title_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        worksheet.row_dimensions[anchor_row].height = 24

        if chart_path:
            try:
                image = ExcelImage(str(chart_path))
                image.width = 880
                image.height = 440
                worksheet.add_image(image, f"A{anchor_row + 3}")
            except Exception:
                worksheet.cell(row=anchor_row + 3, column=1, value="Chart image could not be embedded.")


def _relative_path(path: Any, package_dir: Path) -> str:
    if path in (None, "") or pd.isna(path):
        return ""
    raw = Path(str(path))
    try:
        return raw.relative_to(package_dir).as_posix()
    except ValueError:
        return raw.as_posix()


def _short_label(value: str, *, max_length: int) -> str:
    text = str(value or "").strip()
    if len(text) <= max_length:
        return text
    return text[: max_length - 1].rstrip() + "..."


def _safe_date(value: Any) -> str | None:
    if value is None or pd.isna(value):
        return None
    return str(pd.to_datetime(value).date())


def _is_all_token(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return value.strip().lower() in {"all", "*", "any"}
    return False


def _split_list(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, str):
        return [item.strip() for item in value.split(",") if item.strip()]
    if isinstance(value, (list, tuple, set)):
        return list(value)
    return [value]


def _as_str_list(value: Any) -> list[str]:
    return [str(item).strip() for item in _split_list(value) if str(item).strip()]


def _as_float_list(value: Any) -> list[float]:
    return [float(item) for item in _split_list(value) if str(item).strip()]


def _as_int_list(value: Any) -> list[int]:
    return [int(float(item)) for item in _split_list(value) if str(item).strip()]


def _as_optional_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    return float(value)


def _as_optional_int(value: Any) -> int | None:
    if value in (None, ""):
        return None
    return int(float(value))
