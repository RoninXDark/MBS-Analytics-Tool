from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from bond_story_reporting.data import normalize_bloomberg_wide_frame
from bond_story_reporting.reporting import build_idea_chart_package, generate_ranked_ideas


def _normalized_sample() -> pd.DataFrame:
    wide = pd.DataFrame(
        {
            "Story": [
                "2023 FNCL 5.0 NY LB125",
                "2023 FNCL 5.0 FL LB125",
                "2025 GNMA 4.5 TX",
            ],
            "2018-12-28": [100.0, 101.0, 102.0],
            "2019-01-04": [100.0, 101.0, 102.0],
            "2019-01-11": [99.0, 103.0, 101.0],
            "2019-01-18": [98.0, 105.0, 100.0],
        }
    )
    normalized, _ = normalize_bloomberg_wide_frame(wide)
    return normalized


def test_build_idea_chart_package_writes_expected_outputs(tmp_path: Path):
    normalized = _normalized_sample()
    ranking = generate_ranked_ideas(normalized, ranking_basis="z_score", top_n=2, min_observations=2)

    result = build_idea_chart_package(
        normalized,
        ranking,
        tmp_path / "ideas",
        top_n=1,
        include_full_history=True,
        value_label="Payup",
    )

    assert result.buy_candidates_path.exists()
    assert result.sell_candidates_path.exists()
    assert result.readable_summary_path is not None
    assert result.readable_summary_path.exists()
    assert result.readable_workbook_path is not None
    assert result.readable_workbook_path.exists()
    assert result.readable_workbook_path.name == "CLIENT_REPORT_BUY_SELL.xlsx"
    assert not (result.output_dir / "buy_sell_summary.xlsx").exists()
    assert result.buy_candidates_path.parent.name == "diagnostics"
    assert result.sell_candidates_path.parent.name == "diagnostics"
    assert result.chart_metadata_path.exists()
    assert result.package_metadata_path.exists()
    assert (result.output_dir / "charts" / "buy").exists()
    assert (result.output_dir / "charts" / "sell").exists()
    assert len(list((result.output_dir / "charts" / "buy").glob("*.png"))) == 1
    assert len(list((result.output_dir / "charts" / "sell").glob("*.png"))) == 1
    assert result.metadata["charts_generated"] == 2
    assert result.metadata["buy_candidates"] == 1
    assert result.metadata["sell_candidates"] == 1
    assert result.metadata["client_readable_outputs"]["client_report_xlsx"]
    assert result.metadata["client_readable_outputs"]["technical_summary_csv"]
    assert result.readable_summary_path.read_text(encoding="utf-8-sig").startswith("sep=,")

    workbook = load_workbook(result.readable_workbook_path)
    assert workbook.sheetnames == [
        "Run Summary",
        "Buy Ideas",
        "Sell Ideas",
        "Technical Metadata",
        "Buy Charts",
        "Sell Charts",
    ]
    for sheet_name in ["Buy Ideas", "Sell Ideas"]:
        worksheet = workbook[sheet_name]
        headers = [cell.value for cell in worksheet[1]]
        assert "Rank" in headers
        assert "Story Name" in headers
        assert "View in Workbook" in headers
        assert "Chart File" in headers
        assert "Story / Attribute" in headers
        assert "Geography" not in headers
        assert "Loan Balance Bucket" not in headers
        assert "Story Family" not in headers
        assert "Story Tag" not in headers
        assert "comparable_key" not in headers
        assert "Side" not in headers
        assert worksheet.freeze_panes == "A2"
        assert worksheet.auto_filter.ref is not None

    buy_sheet = workbook["Buy Ideas"]
    buy_headers = [cell.value for cell in buy_sheet[1]]
    coupon_col = buy_headers.index("Coupon") + 1
    year_col = buy_headers.index("Vintage") + 1
    attribute_col = buy_headers.index("Story / Attribute") + 1
    story_col = buy_headers.index("Story Name") + 1
    embedded_col = buy_headers.index("View in Workbook") + 1
    chart_col = buy_headers.index("Chart File") + 1
    assert buy_sheet.cell(row=2, column=coupon_col).data_type == "s"
    assert buy_sheet.cell(row=2, column=coupon_col).value in {"5.0", "4.5"}
    assert isinstance(buy_sheet.cell(row=2, column=year_col).value, int)
    assert buy_sheet.cell(row=2, column=attribute_col).value in {"NY / LB125", "FL / LB125", "TX"}
    buy_chart_cell = buy_sheet.cell(row=2, column=chart_col)
    buy_story_cell = buy_sheet.cell(row=2, column=story_col)
    buy_embedded_cell = buy_sheet.cell(row=2, column=embedded_col)
    assert buy_chart_cell.value.startswith("charts/buy/")
    assert buy_chart_cell.hyperlink is not None
    assert buy_story_cell.hyperlink is not None
    assert buy_embedded_cell.hyperlink is not None
    assert buy_chart_cell.hyperlink.target == buy_chart_cell.value
    assert buy_story_cell.hyperlink.target == buy_chart_cell.value
    assert buy_embedded_cell.hyperlink.target.startswith("#'Buy Charts'!A")
    assert buy_chart_cell.font.underline == "single"
    assert buy_story_cell.font.underline == "single"
    assert buy_embedded_cell.font.underline == "single"
    assert len(workbook["Buy Charts"]._images) == 1

    sell_sheet = workbook["Sell Ideas"]
    sell_headers = [cell.value for cell in sell_sheet[1]]
    sell_story_col = sell_headers.index("Story Name") + 1
    sell_embedded_col = sell_headers.index("View in Workbook") + 1
    sell_chart_col = sell_headers.index("Chart File") + 1
    sell_chart_cell = sell_sheet.cell(row=2, column=sell_chart_col)
    sell_story_cell = sell_sheet.cell(row=2, column=sell_story_col)
    sell_embedded_cell = sell_sheet.cell(row=2, column=sell_embedded_col)
    assert sell_chart_cell.value.startswith("charts/sell/")
    assert sell_chart_cell.hyperlink is not None
    assert sell_story_cell.hyperlink is not None
    assert sell_embedded_cell.hyperlink is not None
    assert sell_chart_cell.hyperlink.target == sell_chart_cell.value
    assert sell_story_cell.hyperlink.target == sell_chart_cell.value
    assert sell_embedded_cell.hyperlink.target.startswith("#'Sell Charts'!A")
    assert len(workbook["Sell Charts"]._images) == 1

    technical_headers = [cell.value for cell in workbook["Technical Metadata"][1]]
    assert "comparable_key" in technical_headers
    technical_chart_col = technical_headers.index("chart_path") + 1
    technical_chart_cell = workbook["Technical Metadata"].cell(row=2, column=technical_chart_col)
    assert technical_chart_cell.value.startswith("charts/")
    assert technical_chart_cell.hyperlink is not None
    assert technical_chart_cell.hyperlink.target == technical_chart_cell.value
    assert technical_chart_cell.font.underline == "single"


def test_build_idea_chart_package_can_limit_chart_history(tmp_path: Path):
    normalized = _normalized_sample()
    ranking = generate_ranked_ideas(normalized, ranking_basis="payup", top_n=1, min_observations=2)

    result = build_idea_chart_package(
        normalized,
        ranking,
        tmp_path / "ideas_active",
        top_n=1,
        include_full_history=False,
    )

    assert result.metadata["include_full_history"] is False
    assert not result.chart_metadata.empty
    # Active focus is based on origination year, not observation date. These
    # sample stories are active vintages, so all four observations remain.
    assert result.chart_metadata["series_points"].min() == 4
