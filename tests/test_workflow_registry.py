import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from bond_story_reporting.automation import (
    UnsupportedWorkflowError,
    available_workflows,
    resolve_saved_workflow,
    run_saved_batch,
)
from bond_story_reporting.reporting.payup_curve_package import _coupon_axis_ticks


def _write_wide_input(path: Path) -> None:
    pd.DataFrame(
        {
            "Story": [
                "2023 FNCL 5.0 NY LB125",
                "2023 FNCL 5.0 FL LB125",
                "2025 GNMA 4.5 TX",
            ],
            "2019-01-04": [100.0, 101.0, 102.0],
            "2019-01-11": [99.0, 103.0, 101.0],
            "2019-01-18": [98.0, 105.0, 100.0],
        }
    ).to_csv(path, index=False)


def _write_store(path: Path, configs: dict) -> None:
    path.write_text(json.dumps({"configs": configs}), encoding="utf-8")


def test_workflow_registry_exposes_current_and_future_workflows():
    workflows = available_workflows()
    assert "idea_scan" in workflows
    assert "weekly_report" in workflows
    assert "payup_curve_pack" in workflows
    assert "report_suite" in workflows


def test_run_config_routes_existing_idea_scan_config(tmp_path: Path):
    input_path = tmp_path / "wide.csv"
    _write_wide_input(input_path)
    store = tmp_path / "configs.json"
    _write_store(
        store,
        {
            "Existing Ideas": {
                "workflow": "idea_scan",
                "input_path": "wide.csv",
                "output_dir": "overnight",
                "ranking_method": "z_score",
                "top_n": 1,
                "package_name": "registry_existing",
                "min_observations": 2,
            }
        },
    )

    result = run_saved_batch("Existing Ideas", config_store_path=store, base_dir=tmp_path)

    assert result["workflow"] == "idea_scan"
    assert result["config_name"] == "Existing Ideas"
    assert result["metadata"]["ranking_strategy"]["name"] == "z_score"
    assert result["metadata"]["client_summary"]["summary_workbook"].endswith("CLIENT_REPORT_BUY_SELL.xlsx")
    assert (tmp_path / "overnight" / "registry_existing" / "diagnostics" / "run_metadata.json").exists()


def test_run_config_missing_workflow_defaults_to_idea_scan_for_idea_scan_shape(tmp_path: Path):
    input_path = tmp_path / "wide.csv"
    _write_wide_input(input_path)
    store = tmp_path / "configs.json"
    _write_store(
        store,
        {
            "No Workflow Ideas": {
                "input_path": "wide.csv",
                "output_dir": "overnight",
                "ranking_method": "payup",
                "top_n": 1,
                "package_name": "registry_inferred",
                "min_observations": 2,
            }
        },
    )

    result = run_saved_batch("No Workflow Ideas", config_store_path=store, base_dir=tmp_path)

    assert result["workflow"] == "idea_scan"
    assert result["metadata"]["ranking_strategy"]["name"] == "payup"
    assert (tmp_path / "overnight" / "registry_inferred" / "CLIENT_REPORT_BUY_SELL.xlsx").exists()


def test_legacy_weekly_shape_still_resolves_to_weekly_report():
    resolution = resolve_saved_workflow({"auto_input": "sample.csv", "auto_output": "out"})
    assert resolution.workflow == "weekly_report"
    assert resolution.inferred is True


def test_payup_curve_pack_runs_through_workflow_registry(tmp_path: Path):
    input_path = tmp_path / "wide.csv"
    _write_wide_input(input_path)
    store = tmp_path / "configs.json"
    _write_store(
        store,
        {
            "Payup Curves": {
                "workflow": "payup_curve_pack",
                "input_path": "wide.csv",
                "output_dir": "canned",
                "report_name": "Nightly Payup Curves",
                "package_name": "payup_registry",
                "chart_requests": [
                    {
                        "name": "LB125 2023 Time Series",
                        "chart_type": "time_series",
                        "loan_balance_buckets": ["LB125"],
                        "origination_years": [2023],
                        "coupons": "all",
                    },
                    {
                        "name": "Coupon Snapshot",
                        "chart_type": "latest_snapshot_curve",
                        "programs": ["FNCL"],
                        "origination_years": [2023],
                        "coupon_min": 4.5,
                        "coupon_max": 5.5,
                    },
                ],
            }
        },
    )

    result = run_saved_batch("Payup Curves", config_store_path=store, base_dir=tmp_path)

    package_dir = tmp_path / "canned" / "payup_registry"
    assert result["workflow"] == "payup_curve_pack"
    assert result["metadata"]["client_summary"]["charts_generated"] == 2
    workbook_path = package_dir / "CLIENT_REPORT_INDEX.xlsx"
    assert workbook_path.exists()
    assert len(list((package_dir / "charts").glob("*.png"))) == 2
    assert (package_dir / "diagnostics" / "run_metadata.json").exists()
    assert (package_dir / "diagnostics" / "chart_metadata.csv").exists()
    assert (package_dir / "diagnostics" / "parse_diagnostics.json").exists()

    workbook = load_workbook(workbook_path)
    chart_sheet = workbook["Generated Charts"]
    headers = [cell.value for cell in chart_sheet[1]]
    request_name_col = headers.index("Request Name") + 1
    chart_title_col = headers.index("Chart Title") + 1
    attribute_col = headers.index("Story / Attribute") + 1
    x_axis_col = headers.index("X-Axis") + 1
    embedded_col = headers.index("View in Workbook") + 1
    chart_file_col = headers.index("Chart File") + 1
    chart_file_cell = chart_sheet.cell(row=2, column=chart_file_col)
    request_name_cell = chart_sheet.cell(row=2, column=request_name_col)
    embedded_cell = chart_sheet.cell(row=2, column=embedded_col)
    chart_title_cell = chart_sheet.cell(row=2, column=chart_title_col)
    attribute_cell = chart_sheet.cell(row=2, column=attribute_col)
    snapshot_x_axis_cell = chart_sheet.cell(row=3, column=x_axis_col)
    assert attribute_cell.value == "LB125"
    assert "FNCL" in chart_title_cell.value
    assert "2023 Vintage" in chart_title_cell.value
    assert chart_file_cell.value.startswith("charts/")
    assert chart_file_cell.hyperlink is not None
    assert request_name_cell.hyperlink is not None
    assert embedded_cell.hyperlink is not None
    assert chart_file_cell.hyperlink.target == chart_file_cell.value
    assert request_name_cell.hyperlink.target == chart_file_cell.value
    assert embedded_cell.hyperlink.target.startswith("#'Charts'!A")
    assert chart_file_cell.font.underline == "single"
    assert request_name_cell.font.underline == "single"
    assert embedded_cell.font.underline == "single"
    assert "Charts" in workbook.sheetnames
    assert len(workbook["Charts"]._images) == 2
    assert snapshot_x_axis_cell.value == "coupon"
    assert _coupon_axis_ticks([4.5, 5.0, 5.5], {"coupon_min": 4.5, "coupon_max": 5.5}) == [4.5, 5.0, 5.5]


def test_payup_snapshot_uses_vintage_axis_for_single_coupon_multi_vintage(tmp_path: Path):
    input_path = tmp_path / "wide.csv"
    pd.DataFrame(
        {
            "Story": [
                "2022 FNCL 5.0 NY",
                "2023 FNCL 5.0 NY",
                "2024 FNCL 5.0 NY",
            ],
            "2019-01-04": [100.0, 101.0, 102.0],
            "2019-01-11": [101.0, 102.0, 103.0],
            "2019-01-18": [102.0, 103.0, 104.0],
        }
    ).to_csv(input_path, index=False)
    store = tmp_path / "configs.json"
    _write_store(
        store,
        {
            "Vintage Curve": {
                "workflow": "payup_curve_pack",
                "input_path": "wide.csv",
                "output_dir": "canned",
                "report_name": "Vintage Curve",
                "package_name": "vintage_curve",
                "chart_requests": [
                    {
                        "name": "FNCL NY 5.0 Vintage Curve",
                        "chart_type": "latest_snapshot_curve",
                        "programs": ["FNCL"],
                        "geographies": ["NY"],
                        "coupons": [5.0],
                        "origination_year_min": 2022,
                        "origination_year_max": 2024,
                    }
                ],
            }
        },
    )

    result = run_saved_batch("Vintage Curve", config_store_path=store, base_dir=tmp_path)

    assert result["metadata"]["client_summary"]["charts_generated"] == 1
    workbook = load_workbook(tmp_path / "canned" / "vintage_curve" / "CLIENT_REPORT_INDEX.xlsx")
    chart_sheet = workbook["Generated Charts"]
    headers = [cell.value for cell in chart_sheet[1]]
    assert chart_sheet.cell(row=2, column=headers.index("X-Axis") + 1).value == "vintage"


def test_payup_snapshot_skips_ambiguous_coupon_and_vintage_mix(tmp_path: Path):
    input_path = tmp_path / "wide.csv"
    pd.DataFrame(
        {
            "Story": [
                "2022 FNCL 4.5 NY",
                "2023 FNCL 5.0 NY",
            ],
            "2019-01-04": [100.0, 101.0],
            "2019-01-11": [101.0, 102.0],
            "2019-01-18": [102.0, 103.0],
        }
    ).to_csv(input_path, index=False)
    store = tmp_path / "configs.json"
    _write_store(
        store,
        {
            "Ambiguous Curve": {
                "workflow": "payup_curve_pack",
                "input_path": "wide.csv",
                "output_dir": "canned",
                "report_name": "Ambiguous Curve",
                "package_name": "ambiguous_curve",
                "chart_requests": [
                    {
                        "name": "Ambiguous NY Curve",
                        "chart_type": "latest_snapshot_curve",
                        "programs": ["FNCL"],
                        "geographies": ["NY"],
                    }
                ],
            }
        },
    )

    result = run_saved_batch("Ambiguous Curve", config_store_path=store, base_dir=tmp_path)

    assert result["metadata"]["client_summary"]["charts_generated"] == 0
    assert result["metadata"]["client_summary"]["skipped_requests"] == 1
    workbook = load_workbook(tmp_path / "canned" / "ambiguous_curve" / "CLIENT_REPORT_INDEX.xlsx")
    summary_sheet = workbook["Chart Requests Summary"]
    headers = [cell.value for cell in summary_sheet[1]]
    status_col = headers.index("Status") + 1
    assert summary_sheet.cell(row=2, column=status_col).value == "ambiguous_snapshot_axes"


def test_payup_curve_pack_empty_request_is_reported_without_crashing(tmp_path: Path):
    input_path = tmp_path / "wide.csv"
    _write_wide_input(input_path)
    store = tmp_path / "configs.json"
    _write_store(
        store,
        {
            "Empty Payup": {
                "workflow": "payup_curve_pack",
                "input_path": "wide.csv",
                "output_dir": "canned",
                "report_name": "Empty Payup",
                "package_name": "empty_payup",
                "chart_requests": [
                    {
                        "name": "No Match",
                        "chart_type": "time_series",
                        "geographies": ["ZZ"],
                    }
                ],
            }
        },
    )

    result = run_saved_batch("Empty Payup", config_store_path=store, base_dir=tmp_path)

    package_dir = tmp_path / "canned" / "empty_payup"
    assert result["workflow"] == "payup_curve_pack"
    assert result["metadata"]["client_summary"]["charts_generated"] == 0
    assert result["metadata"]["client_summary"]["skipped_requests"] == 1
    assert (package_dir / "CLIENT_REPORT_INDEX.xlsx").exists()
    assert (package_dir / "diagnostics" / "run_metadata.json").exists()


def test_payup_curve_pack_requires_explicit_filter_fields(tmp_path: Path):
    input_path = tmp_path / "wide.csv"
    _write_wide_input(input_path)
    store = tmp_path / "configs.json"
    _write_store(
        store,
        {
            "Ambiguous Payup": {
                "workflow": "payup_curve_pack",
                "input_path": "wide.csv",
                "output_dir": "canned",
                "report_name": "Ambiguous Payup",
                "chart_requests": [
                    {
                        "name": "LB125 2023 Time Series",
                        "chart_type": "time_series",
                    }
                ],
            }
        },
    )

    with pytest.raises(ValueError, match="missing explicit filter criteria"):
        run_saved_batch("Ambiguous Payup", config_store_path=store, base_dir=tmp_path)


def test_unsupported_future_workflow_fails_clearly(tmp_path: Path):
    store = tmp_path / "configs.json"
    _write_store(
        store,
        {
            "Future Suite": {"workflow": "report_suite"}
        },
    )

    with pytest.raises(UnsupportedWorkflowError, match="reserved for a future phase"):
        run_saved_batch("Future Suite", config_store_path=store, base_dir=tmp_path)


def test_unknown_workflow_fails_with_supported_workflow_list(tmp_path: Path):
    store = tmp_path / "configs.json"
    _write_store(store, {"Mystery": {"workflow": "unknown_magic_workflow"}})

    with pytest.raises(UnsupportedWorkflowError, match="Currently implemented workflows"):
        run_saved_batch("Mystery", config_store_path=store, base_dir=tmp_path)
